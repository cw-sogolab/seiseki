#!/usr/bin/env python3
"""월간 수업계획서 만들기

지난달 계획서(워드) 폴더를 양식으로 삼아, 연간 계획표의 개강·종강·휴강일로 다음 달 날짜·회차를 채웁니다.

    python3 class_plan.py --month 2026-11 --prev "…/강의계획서/26/10" --out "…/강의계획서/26/11"
    python3 class_plan.py --check            # 연간 계획표 횟수와 계산이 맞는지만 확인

- 반의 요일(토일·화목 등)은 지난달 파일의 날짜에서 알아냅니다.
- 개념·단풀: 마지막 회차 = 월말테스트. 개념: 매주 마지막 수업일과 월말테스트일 과제 = 주간 숙제.
- 수업 내용·교재 범위는 비워 두고, 요약 엑셀에서 노란 칸으로 표시합니다.
- 지난달 파일이 이번 달 첫날까지 잡혀 있으면(예: 10월 실전D의 11/1) 그 날짜는 빼고 알려 줍니다.
- 이미 있는 파일은 덮어쓰지 않습니다 (--force 로 덮어쓰기).
필요: pip3 install python-docx openpyxl
"""
import argparse, json, re, sys, unicodedata
from datetime import date, timedelta
from pathlib import Path

WD = "월화수목금토일"
PATTERNS = {"월수금": {0, 2, 4}, "화목토": {1, 3, 5}, "화목": {1, 3}, "토일": {5, 6}}
HERE = Path(__file__).resolve().parent


def D(s):
    y, m, d = map(int, s.split("-"))
    return date(y, m, d)


def label(d):
    return f"{d.month}월{d.day}일"


def load_calendar(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))["months"]


def sessions(cal, ym, weekdays, start=None):
    m = cal[ym]
    s, e = D(m["start"]), D(m["end"])
    if start and start > s:
        s = start
    off = {D(x[0]) for x in m["off"]}
    out, d = [], s
    while d <= e:
        if d.weekday() in weekdays and d not in off:
            out.append(d)
        d += timedelta(days=1)
    return out


def check_calendar(cal):
    bad = []
    for ym, m in cal.items():
        for name, wds in PATTERNS.items():
            got, exp = len(sessions(cal, ym, wds)), m["counts"].get(name)
            if exp is not None and got != exp:
                bad.append(f"{ym} {name}: 계획표 {exp}회, 계산 {got}회")
    return bad


def kind_of(name):
    for k in ("개념", "단풀", "실전"):
        if k in name:
            return k
    return "기타"


def homework_flags(days, kind):
    if kind != "개념":
        return [False] * len(days)
    return [i == len(days) - 1 or days[i + 1].isocalendar()[:2] != d.isocalendar()[:2] for i, d in enumerate(days)]


def plan_rows(days, kind):
    hw = homework_flags(days, kind)
    rows = []
    for i, d in enumerate(days):
        last = i == len(days) - 1
        rows.append({"no": f"{i + 1:02d}", "date": d, "content": "월말테스트" if last and kind in ("개념", "단풀") else "",
                     "range": "", "hw": "주간 숙제" if hw[i] else ""})
    return rows


def uniq_cells(row):
    seen, out = set(), []
    for c in row.cells:
        if id(c._tc) in seen:
            continue
        seen.add(id(c._tc))
        out.append(c)
    return out


def set_text(cell, text):
    p = cell.paragraphs[0]
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ""
    elif text:
        p.add_run(text)


def find_schedule(table):
    """'No.' 머리줄 다음부터가 회차 줄"""
    for i, r in enumerate(table.rows):
        cells = uniq_cells(r)
        if cells and cells[0].text.strip().startswith("No"):
            return i + 1
    raise ValueError("세부 일정 표(No. 줄)를 찾지 못했어요")


def read_prev(doc, prev_ym):
    t = doc.tables[0]
    first = find_schedule(t)
    py, pm = map(int, prev_ym.split("-"))
    dates = []
    for r in t.rows[first:]:
        cells = uniq_cells(r)
        if len(cells) < 2:
            continue
        m = re.search(r"(\d{1,2})\s*[월/.]\s*(\d{1,2})", cells[1].text)
        if not m:
            continue
        mo, da = int(m.group(1)), int(m.group(2))
        y = py + 1 if mo < pm - 6 else (py - 1 if mo > pm + 6 else py)
        dates.append(date(y, mo, da))
    return t, first, dates


def infer_pattern(dates):
    wds = {d.weekday() for d in dates}
    for name, s in PATTERNS.items():
        if wds == s:
            return name, s, ""
    cands = sorted([(len(s), n, s) for n, s in PATTERNS.items() if wds <= s])
    if cands:
        return cands[0][1], cands[0][2], f"요일이 {''.join(WD[w] for w in sorted(wds))}만 있어서 {cands[0][1]}로 봤어요"
    return "".join(WD[w] for w in sorted(wds)), wds, "알려진 요일 묶음이 아니라 지난달 요일 그대로 썼어요"


def build(args):
    import docx
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    cal = load_calendar(args.calendar)
    ym = args.month
    if ym not in cal:
        sys.exit(f"{ym} 이 연간 계획표({args.calendar})에 없어요. 계획표 JSON에 그 달을 먼저 넣어 주세요.")
    bad = check_calendar({ym: cal[ym]})
    y, m = map(int, ym.split("-"))
    prev_ym = f"{y - 1}-12" if m == 1 else f"{y}-{m - 1:02d}"
    prev_dir, out_dir = Path(args.prev), Path(args.out)
    out_dir.mkdir(parents=True, exist_ok=True)
    yymm_prev, yymm = prev_ym[2:4] + prev_ym[5:7], ym[2:4] + ym[5:7]
    month_start = D(cal[ym]["start"])

    report, made = [], []
    srcs = sorted(p for p in prev_dir.glob("*.docx") if not p.name.startswith("~$"))
    if not srcs:
        sys.exit(f"{prev_dir} 에 지난달 계획서(.docx)가 없어요.")
    for src in srcs:
        name = unicodedata.normalize("NFC", src.name)
        if args.only and not any(o in name for o in args.only):
            continue
        doc = docx.Document(str(src))
        t, first, dates = read_prev(doc, prev_ym)
        if not dates:
            report.append(f"- {name}: 날짜를 못 읽어서 건너뜀")
            continue
        pname, wds, note = infer_pattern(dates)
        kind = kind_of(name)
        start = None
        if max(dates) >= month_start:
            start = max(dates) + timedelta(days=1)
            note = (note + " / " if note else "") + f"지난달 계획서에 {label(max(dates))}까지 있어서 그 다음 날부터 잡았어요"
        days = sessions(cal, ym, wds, start)
        rows = plan_rows(days, kind)
        slots = len(t.rows) - first
        if len(rows) > slots:
            note = (note + " / " if note else "") + f"회차 {len(rows)}개가 표 칸 {slots}줄보다 많아요 — 워드 표에 줄을 추가해야 해요"
        for p in doc.paragraphs:
            if f"{prev_ym[5:7].lstrip('0')}월강의계획" in p.text.replace(" ", "") and p.runs:
                txt = re.sub(r"(\d{4})\s*년\s*\d{1,2}\s*월", f"{y}년{m}월", p.text)
                p.runs[0].text = txt
                for r in p.runs[1:]:
                    r.text = ""
        for i in range(slots):
            cells = uniq_cells(t.rows[first + i])
            if len(cells) < 5:
                continue
            if i < len(rows):
                r = rows[i]
                vals = [label(r["date"]), r["content"], r["range"], r["hw"]]
            else:
                vals = ["", "", "", ""]
            for k, v in enumerate(vals, start=1):
                set_text(cells[k], v)
        out_name = (yymm + name[len(yymm_prev):]) if name.startswith(yymm_prev) else f"{yymm}_{name}"
        dest = out_dir / out_name
        if dest.exists() and not args.force:
            report.append(f"- {out_name}: 이미 있어서 건너뜀 (--force 로 덮어쓰기)")
            continue
        doc.save(str(dest))
        title = doc.paragraphs[1].text.strip().strip("<> ").strip() if len(doc.paragraphs) > 1 else name
        header = []
        for r in t.rows[:first - 2]:
            vals = [c.text.strip() for c in uniq_cells(r) if c.text.strip()]
            if vals:
                header.append(vals)
        made.append((out_name, title, pname, kind, rows, header))
        report.append(f"- {out_name}: {pname} {len(rows)}회 ({', '.join(label(r['date']) for r in rows)})" + (f"  ※ {note}" if note else ""))

    # 요약 엑셀
    thin = Side(style="thin", color="999999")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    navy = PatternFill("solid", fgColor="12233F")
    need = PatternFill("solid", fgColor="FFF2CC")
    wb = Workbook()
    ws = wb.active
    ws.title = f"{m}월 회차표"
    mm = cal[ym]
    ws["A1"] = f"{y}년 {m}월 요일별 수업일"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = (f"개강 {D(mm['start']).month}/{D(mm['start']).day} · 종강 {D(mm['end']).month}/{D(mm['end']).day} · 휴강 "
                + (", ".join(f"{D(x).month}/{D(x).day}({r})" for x, r in mm["off"]) or "없음"))
    col = 1
    for pname, wds in PATTERNS.items():
        days = sessions(cal, ym, wds)
        for k, v in ((0, pname), (1, f"{len(days)}회")):
            c = ws.cell(4, col + k, v)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = navy
        for i, d in enumerate(days):
            ws.cell(5 + i, col, f"{i + 1:02d}").border = box
            ws.cell(5 + i, col + 1, f"{label(d)}({WD[d.weekday()]})").border = box
        ws.column_dimensions[ws.cell(1, col).column_letter].width = 7
        ws.column_dimensions[ws.cell(1, col + 1).column_letter].width = 15
        ws.column_dimensions[ws.cell(1, col + 2).column_letter].width = 3
        col += 3
    for out_name, title, pname, kind, rows, header in made:
        sh = wb.create_sheet(re.sub(r"[\\/?*\[\]:]", " ", Path(out_name).stem)[:31])
        for c, w in zip("ABCDE", (7, 16, 40, 16, 13)):
            sh.column_dimensions[c].width = w
        sh["A1"] = f"{y}년 {m}월 강의계획"
        sh["A1"].font = Font(bold=True, size=15)
        sh["A2"] = f"< {title} >  ({pname} · {kind})"
        sh["A2"].font = Font(bold=True, size=12)
        r0 = 4
        for vals in header:
            sh.cell(r0, 1, vals[0]).font = Font(bold=True)
            sh.cell(r0, 2, "  ".join(vals[1:]))
            sh.merge_cells(start_row=r0, start_column=2, end_row=r0, end_column=5)
            r0 += 1
        r0 += 1
        for j, h in enumerate(["No.", "날 짜 (요일)", "수 업 내 용", "주 교재 범위", "과 제"], start=1):
            c = sh.cell(r0, j, h)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = navy
            c.border = box
            c.alignment = Alignment(horizontal="center")
        for i, r in enumerate(rows):
            vals = [r["no"], f"{label(r['date'])}({WD[r['date'].weekday()]})", r["content"], r["range"], r["hw"]]
            for j, v in enumerate(vals, start=1):
                c = sh.cell(r0 + 1 + i, j, v)
                c.border = box
                if j in (3, 4) and not v:
                    c.fill = need
        sh.cell(r0 + len(rows) + 2, 1, "노란 칸 = 이번 달 진도(강·교재 범위)를 채울 곳").font = Font(italic=True, color="7F6000")
    v = wb.create_sheet("검증")
    v["A1"] = "연간 계획표 횟수와 계산 대조"
    v["A1"].font = Font(bold=True)
    v.append([])
    v.append(["월"] + [f"{n}(계획표/계산)" for n in PATTERNS] + ["결과"])
    for k, mm2 in cal.items():
        got = [len(sessions(cal, k, s)) for s in PATTERNS.values()]
        exp = [mm2["counts"].get(n) for n in PATTERNS]
        v.append([k] + [f"{e}/{g}" for e, g in zip(exp, got)] + ["일치" if exp == got else "불일치"])
    xlsx = out_dir / f"{yymm}_수업계획서.xlsx"
    if xlsx.exists() and not args.force:
        xlsx = out_dir / f"{yymm}_수업계획서_새로만듦.xlsx"
    wb.save(str(xlsx))
    print(f"{ym} 수업계획서 — 지난달 {len(srcs)}개 파일 기준")
    if bad:
        print("⚠ 연간 계획표 횟수와 계산이 달라요: " + "; ".join(bad))
    print("\n".join(report))
    print(f"요약 엑셀: {xlsx}")


def main():
    ap = argparse.ArgumentParser(description="월간 수업계획서 만들기")
    ap.add_argument("--month", help="만들 달 YYYY-MM")
    ap.add_argument("--prev", help="지난달 계획서(.docx) 폴더")
    ap.add_argument("--out", help="저장할 폴더")
    ap.add_argument("--calendar", default=str(HERE / "calendar_2026.json"))
    ap.add_argument("--only", nargs="*", help="파일 이름에 이 글자가 든 것만")
    ap.add_argument("--force", action="store_true", help="이미 있는 파일 덮어쓰기")
    ap.add_argument("--check", action="store_true", help="연간 계획표 검증만")
    a = ap.parse_args()
    if a.check:
        bad = check_calendar(load_calendar(a.calendar))
        print("전부 일치" if not bad else "\n".join(bad))
        sys.exit(1 if bad else 0)
    if not (a.month and a.prev and a.out):
        ap.error("--month, --prev, --out 이 필요해요")
    build(a)


if __name__ == "__main__":
    main()
